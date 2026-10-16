"""Figures, tables and the Excel model. Everything is generated from pipeline outputs (no manual numbers).

Visual conventions (dataviz reference palette, light mode): categorical slots in fixed order,
2px lines, recessive grid, no dual axes, legend + selective direct labels, diverging blue<->red with a
gray midpoint for P&L heatmaps.
"""
from __future__ import annotations

import json
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
from matplotlib.colors import LinearSegmentedColormap, TwoSlopeNorm  # noqa: E402

SERIES = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300", "#4a3aa7", "#e34948"]
SURFACE, INK, INK2, GRID = "#fcfcfb", "#0b0b0b", "#52514e", "#e4e3df"
DIVERGING = LinearSegmentedColormap.from_list("div", ["#e34948", "#f0efec", "#2a78d6"])

plt.rcParams.update({
    "figure.facecolor": SURFACE, "axes.facecolor": SURFACE, "savefig.facecolor": SURFACE,
    "axes.edgecolor": GRID, "axes.labelcolor": INK2, "xtick.color": INK2, "ytick.color": INK2,
    "axes.grid": True, "grid.color": GRID, "grid.linewidth": 0.6, "axes.spines.top": False,
    "axes.spines.right": False, "lines.linewidth": 2.0, "font.size": 9.5, "axes.titlesize": 11,
    "axes.titleweight": "bold", "axes.titlecolor": INK, "legend.frameon": False, "figure.dpi": 110,
})


def _save(fig, ctx, name: str) -> None:
    fig.tight_layout()
    fig.savefig(ctx.fig / name, dpi=200)
    plt.close(fig)


def _src(ax, text: str, y: float = -0.16) -> None:
    ax.annotate(text, xy=(0, y), xycoords="axes fraction", fontsize=7.5, color=INK2, ha="left", va="top")


# ------------------------------------------------------------------ figures
def fig_smlt(ctx) -> None:
    P = ctx.read("prices.csv", index_col=0, parse_dates=True)
    px = P["SMLT"].dropna()
    dd = px / px.cummax() - 1
    ev = pd.read_csv(ctx.root / "data" / "manual" / "event_chronology.csv")
    fig, (a, b) = plt.subplots(2, 1, figsize=(10, 6.2), sharex=True, gridspec_kw={"height_ratios": [2.2, 1]})
    a.plot(px.index, px.values, color=SERIES[0])
    a.set_yscale("log")
    a.set_title("SMLT close (TQBR, RUB, log scale), 2022-01-03 to 2026-09-25")
    for _, r in ev.iterrows():
        d = pd.to_datetime(str(r["event_date"])[:10], errors="coerce")
        if pd.notna(d) and d >= px.index.min():
            a.axvline(d, color=INK2, lw=0.8, ls=":")
    a.annotate(f"{px.iloc[-1]:.1f}", xy=(px.index[-1], px.iloc[-1]), xytext=(4, 0), textcoords="offset points",
               color=INK, fontsize=8.5, va="center")
    b.fill_between(dd.index, dd.values * 100, 0, color=SERIES[7], alpha=0.25, lw=0)
    b.plot(dd.index, dd.values * 100, color=SERIES[7], lw=1.2)
    b.set_ylabel("drawdown from running peak, %")
    _src(b, "Source: MOEX ISS (raw snapshots, data/source_manifest.csv). Dotted lines: dated events in data/manual/event_chronology.csv")
    _save(fig, ctx, "01_smlt_price_drawdown.png")


def fig_credit(ctx) -> None:
    cs = ctx.read("credit_state.csv", index_col=0, parse_dates=True)
    if cs.dropna(how="all").empty:
        return
    fig, (a, b) = plt.subplots(2, 1, figsize=(10, 5.6), sharex=True)
    z = cs["zspread_median_pp"].dropna()
    a.plot(z.index, z.values, color=SERIES[7], lw=1.4)
    a.set_title("Samolet bonds: median Z-spread over MOEX G-curve (percentage points), traded issues only")
    p = cs["price_median_pct_par"].dropna()
    b.plot(p.index, p.values, color=SERIES[0], lw=1.4)
    b.set_title("Samolet bonds: median clean price, % of par")
    _src(b, "Source: MOEX ISS bond history (ZSPREAD, CLOSE); issues with value >= RUB 1m that day and > 90 days to maturity")
    _save(fig, ctx, "02_credit_stress_state.png")


def fig_rolling(ctx) -> None:
    rh = ctx.read("rolling_hedge_ratios.csv", "out", index_col=0, parse_dates=True)
    fig, (a, b) = plt.subplots(2, 1, figsize=(10, 6.2), sharex=True)
    pairs = [("SMLT_FUT", "SMLT futures"), ("MIX_FUT", "MOEX index futures"), ("PIKK", "PIK (peer)"), ("RGBI_FUT", "RGBI futures")]
    for i, (k, lab) in enumerate(pairs):
        c = f"corr_{k}_120d"
        if c in rh:
            s = rh[c].dropna()
            a.plot(s.index, s.values, color=SERIES[i], lw=1.4, label=lab)
            a.annotate(lab, xy=(s.index[-1], s.iloc[-1]), xytext=(4, 0), textcoords="offset points", fontsize=8, color=INK)
    a.axhline(0, color=INK2, lw=0.8)
    a.set_title("Rolling 120-day correlation with SMLT daily returns")
    a.legend(loc="lower left", ncol=4, fontsize=8)
    for i, (k, lab) in enumerate(pairs[:2]):
        c = f"hstar_{k}_120d"
        if c in rh:
            s = rh[c].dropna()
            b.plot(s.index, s.values, color=SERIES[i], lw=1.4, label=lab)
    b.set_title("Rolling 120-day minimum-variance hedge ratio h* = Cov(S,H)/Var(H)")
    b.legend(loc="upper left", fontsize=8)
    _src(b, "Trailing windows, not shifted (stamp t uses returns <= t). Source: outputs/rolling_hedge_ratios.csv")
    _save(fig, ctx, "03_rolling_corr_hstar.png")


def fig_im(ctx) -> None:
    lim = ctx.read("ncc_futures_limits.csv", parse_dates=["tradedate"])
    fig, a = plt.subplots(figsize=(10, 3.8))
    for i, (asset, lab) in enumerate([("SMLT", "SMLT futures"), ("MIX", "MIX (index)"), ("RGBI", "RGBI futures")]):
        s = lim[lim["assetcode"] == asset].set_index("tradedate")["mr1"].sort_index()
        if len(s):
            a.step(s.index, s.values * 100, where="post", color=SERIES[i], label=lab, lw=1.6)
            a.annotate(f"{lab}: {s.iloc[-1]*100:.0f}%", xy=(s.index[-1], s.iloc[-1] * 100), xytext=(4, 9 - 9 * i),
                       textcoords="offset points", fontsize=8, color=INK)
    a.set_title("NCC base market-risk rate mr1 (initial-margin rate, % of notional), point-in-time")
    a.legend(loc="upper left", fontsize=8)
    _src(a, "Source: MOEX ISS rms/engines/futures/objects/limits (weekly 2022-2026, daily last 90 days)")
    _save(fig, ctx, "04_ncc_im_rates.png")


def fig_capital(ctx) -> None:
    bs = pd.read_csv(ctx.out / "bootstrap_strategy_metrics.csv")
    b = bs[(bs["variant"] == "FHS_FULL")]
    fig, axes = plt.subplots(1, 2, figsize=(11, 4.2))
    for ax, hz in zip(axes, ["6m", "12m"]):
        d = b[b["horizon"] == hz]
        ctrl = d[d["strategy"] == "CONTROL"].set_index("capital")["ES95"]
        for i, strat in enumerate(["FUT_25", "FUT_50", "FUT_75", "FUT_100"]):
            s = d[d["strategy"] == strat].set_index("capital")["ES95"]
            red = (ctrl.reindex(s.index) - s) / ctx.E * 100
            ax.plot(s.index / 1e6, red.values, marker="o", ms=5, color=SERIES[i], label=strat.replace("FUT_", "futures ") + "%")
        ax.set_xlabel("hedge capital, RUB m (IM > 80% of capital: omitted)")
        ax.set_ylabel("ES95 reduction vs CONTROL, % of RUB 10m")
        ax.set_title(f"Capital efficiency, {hz} horizon (FHS bootstrap)")
        ax.legend(fontsize=8)
    _src(axes[0], "Same-capital comparator: CONTROL = RUB 10m SMLT + same hedge capital in the money-market sleeve")
    _save(fig, ctx, "05_capital_efficiency.png")


def fig_complexity(ctx) -> None:
    cx = pd.read_csv(ctx.out / "complexity_frontier.csv")
    d = cx[cx["variant"] == "FHS_FULL"]
    abbrev = {"SMLT_FUT": "F", "INDEX": "I", "SECTOR_PIKK": "S", "RATES_RGBI": "R", "GOLD": "G", "CNY": "C", "OIL_BR": "O"}
    fig, axes = plt.subplots(1, 2, figsize=(11.5, 4.6), sharey=False)
    for ax, cap in zip(axes, [ctx.C, 2_500_000]):
        for i, hz in enumerate(["6m", "12m"]):
            x = d[(d["horizon"] == hz) & (d["capital"] == cap)].sort_values("K")
            if x.empty:
                continue
            ax.plot(x["K"], x["path_ES95"] / 1e6, marker="o", ms=6, color=SERIES[i], label=f"{hz}: full 20k paths (out-of-sample)")
            ax.plot(x["K"], x["lp_ES95_pct"] * ctx.E / 1e6, marker="o", ms=5, ls="--", lw=1.4, color=SERIES[i],
                    alpha=0.6, label=f"{hz}: LP scenario set (in-sample)")
            if hz == "6m":
                for _, r in x.iterrows():
                    ax.annotate("+".join(abbrev.get(a, a) for a in r["active"].split("|")), xy=(r["K"], r["path_ES95"] / 1e6),
                                xytext=(0, 7), textcoords="offset points", fontsize=7, ha="center", color=INK2)
        ax.set_xlabel("K = max number of distinct active sleeves")
        ax.set_ylabel("ES95 of horizon P&L, RUB m (negative = gain)")
        ax.set_title(f"Complexity frontier, hedge capital RUB {cap/1e6:.1f}m")
        ax.legend(fontsize=7, loc="best")
    _src(axes[0], "F=SMLT fut, I=index fut, S=PIK short, R=RGBI fut, G=gold, C=CNY, O=Brent. FHS bootstrap; LP on 4,000 scenarios")
    _save(fig, ctx, "06_complexity_frontier.png")


def fig_tradeoff(ctx) -> None:
    bs = pd.read_csv(ctx.out / "bootstrap_strategy_metrics.csv")
    d = bs[(bs["variant"] == "FHS_FULL") & (bs["capital"] == ctx.C)]
    fig, axes = plt.subplots(1, 2, figsize=(11.5, 4.6))
    colors = {"CONTROL": INK2, "HEDGE_DIRECT": SERIES[0], "DE_RISKING": SERIES[1], "CONVEX_INSURANCE": SERIES[2]}
    for ax, hz in zip(axes, ["6m", "12m"]):
        x = d[d["horizon"] == hz]
        for cls, g in x.groupby("class"):
            mk = "o" if cls != "CONVEX_INSURANCE" else "D"
            ax.scatter(g["ES95"] / ctx.E * 100, g["upside_capture"] * 100, s=34, color=colors.get(cls, SERIES[3]),
                       label=cls + (" (OTC: theoretical)" if cls == "CONVEX_INSURANCE" else ""), marker=mk,
                       edgecolor=SURFACE, linewidth=1)
        for _, r in x.iterrows():
            off = (4, -9) if r["class"] == "DE_RISKING" else (4, 3)
            ax.annotate(r["strategy"].replace("_listedIV", "").replace("_realisedVol", "(RV)").replace("DERISK_SELL_", "SELL_"),
                        xy=(r["ES95"] / ctx.E * 100, r["upside_capture"] * 100), xytext=off,
                        textcoords="offset points", fontsize=6.3, color=INK2)
        ax.set_xlabel("ES95 of horizon P&L, % of RUB 10m (lower = better protection)")
        ax.set_ylabel("upside capture, % of CONTROL gain")
        ax.set_title(f"Protection vs upside retention, {hz}")
    axes[0].legend(fontsize=7.5, loc="lower right")
    _src(axes[0], "FHS bootstrap (volatility rescaled to 2026-09-25), 20,000 paths, seed in configs/project.yaml")
    _save(fig, ctx, "07_protection_vs_upside.png")


def fig_lar(ctx) -> None:
    bs = pd.read_csv(ctx.out / "bootstrap_strategy_metrics.csv")
    d = bs[(bs["variant"] == "FHS_FULL") & (bs["horizon"] == "6m") & bs["strategy"].str.startswith("FUT_")]
    fig, (a, b) = plt.subplots(1, 2, figsize=(11.5, 4.2))
    for i, C in enumerate(sorted(d["capital"].unique())):
        x = d[d["capital"] == C].copy()
        x["h"] = x["strategy"].str.replace("FUT_", "").astype(int)
        x = x.sort_values("h")
        a.plot(x["h"], x["p_sleeve_breach"] * 100, marker="o", ms=5, color=SERIES[i], label=f"hedge capital {C/1e6:.1f}m")
    a.set_xlabel("futures hedge ratio, % (only margin-feasible points)")
    a.set_ylabel("% of paths")
    a.set_title("P(free liquidity sleeve < 0 within 6m)")
    a.legend(fontsize=8)
    f = ctx.interim / "lar_worst_episode_FIXED_100_6m.csv"
    if f.exists():
        s = pd.read_csv(f, index_col=0, parse_dates=True)
        b.plot(s.index, s["vm"].cumsum() / 1e6, color=SERIES[7], lw=1.6, label="cumulative VM")
        b.plot(s.index, s["sleeve_free"] / 1e6, color=SERIES[0], lw=1.6, label="free liquidity sleeve")
        b.axhline(0, color=INK2, lw=0.8)
        b.set_title("Worst historical 6m episode, 100% futures hedge (RUB m)")
        b.legend(fontsize=8)
    _src(a, "FHS bootstrap; IM at NCC mr1 of 2026-09-25 (50%), IM cash earns 0; breach = hedge must be cut or shares sold")
    _save(fig, ctx, "08_liquidity_at_risk.png")


def fig_surface(ctx) -> None:
    sm = pd.read_csv(ctx.out / "scenario_matrix.csv")
    strats = ["CONTROL", "FUT_50", "FUT_100", "INDEX_HSTAR", "OTC_PUT_90_listedIV", "DERISK_SELL_50"]
    strats = [s for s in strats if s in set(sm["strategy"])]
    d = sm[(sm["horizon"] == "12m")]
    vmax = float(np.nanmax(np.abs(d[d["strategy"].isin(strats)]["pnl_pct_of_smlt_10m"]))) * 100
    fig, axes = plt.subplots(len(strats), 3, figsize=(11, 2.0 * len(strats)), sharex=True, sharey=True)
    norm = TwoSlopeNorm(vmin=-vmax, vcenter=0, vmax=vmax)
    for i, st in enumerate(strats):
        for j, rates in enumerate(["DOWN", "STABLE", "UP"]):
            ax = axes[i, j]
            x = d[(d["strategy"] == st) & (d["rates"] == rates)].pivot(index="credit", columns="smlt_terminal_return",
                                                                       values="pnl_pct_of_smlt_10m") * 100
            x = x.reindex(["WORSENING", "STABLE", "IMPROVING"])
            ax.imshow(x.values, cmap=DIVERGING, norm=norm, aspect="auto")
            for (r, c), v in np.ndenumerate(x.values):
                ax.text(c, r, f"{v:.0f}", ha="center", va="center", fontsize=6.5, color=INK)
            ax.set_xticks(range(x.shape[1]), [f"{v:+.0%}" for v in x.columns], fontsize=7)
            ax.set_yticks(range(3), ["worse", "stable", "better"], fontsize=7)
            ax.grid(False)
            if i == 0:
                ax.set_title(f"rates {rates}", fontsize=9)
            if j == 0:
                ax.set_ylabel(st.replace("_listedIV", ""), fontsize=7.5)
    fig.suptitle("12m payoff surface: P&L as % of RUB 10m SMLT, incl. sleeve interest on RUB 10m hedge capital\n"
                 "(x: SMLT terminal return; y: issuer credit state; columns: rates regime; OTC put = theoretical)",
                 fontsize=10.5, fontweight="bold")
    _save(fig, ctx, "09_payoff_surface_12m.png")


def fig_walkforward(ctx) -> None:
    cum = ctx.read("walkforward_cum_pnl.csv", "interim", index_col=0, parse_dates=True)
    keys = [("CONTROL|thr=0.00", "CONTROL (unhedged)"), ("FIXED_50|thr=0.00", "fixed 50%"),
            ("FIXED_100|thr=0.00", "fixed 100%"), ("ROLLING_HSTAR_120|thr=WF_SELECTED", "rolling h*"),
            ("STATE_DEPENDENT|thr=WF_SELECTED", "state-dependent")]
    fig, a = plt.subplots(figsize=(10, 4.2))
    for i, (k, lab) in enumerate(keys):
        if k in cum:
            s = cum[k].dropna() / 1e6
            a.plot(s.index, s.values, color=SERIES[i] if i else INK2, lw=1.6, label=lab)
            a.annotate(lab, xy=(s.index[-1], s.iloc[-1]), xytext=(4, 0), textcoords="offset points", fontsize=7.5, color=INK)
    a.set_title("Walk-forward, constant RUB 10m exposure: cumulative P&L incl. costs and sleeve carry (RUB m)")
    a.legend(fontsize=8, loc="lower left")
    _src(a, "Out-of-sample from 2023-01-03; weekly re-estimation on trailing 120 days; thresholds selected walk-forward")
    _save(fig, ctx, "10_walkforward_cum_pnl.png")


# ------------------------------------------------------------------ Excel model
def excel(ctx) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils.dataframe import dataframe_to_rows

    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    run = json.loads((ctx.out / "run_log.json").read_text()) if (ctx.out / "run_log.json").exists() else {}
    lines = ["SMLT residual-risk hedging model (decision date 2026-09-25)",
             "All sheets except 'Calculator' are pasted values generated by scripts/run_pipeline.py from raw snapshots.",
             "'Calculator' contains live formulas for the first-order payoff algebra (see METHODOLOGY.md).",
             f"Code git HEAD: {run.get('code_git_head', 'n/a')} (run timestamps are in outputs/run_log.json)"]
    for i, l in enumerate(lines, 1):
        ws.cell(row=i, column=1, value=l)
    ws["A1"].font = Font(bold=True, size=12)
    sheets = [("Inputs_at_cutoff", "market_inputs_at_cutoff.csv"), ("Universe", "instrument_universe.csv"),
              ("Candidate_stats", "candidate_stats.csv"), ("Strategy_comparison", "strategy_comparison.csv"),
              ("Risk_metrics", "risk_metrics.csv"), ("Capital_sensitivity", "capital_budget_sensitivity.csv"),
              ("Complexity_frontier", "complexity_frontier.csv"), ("Scenario_matrix", "scenario_matrix.csv"),
              ("Events", "event_overlay.csv"), ("Episodes_LaR", "episodes_summary.csv"),
              ("Options_20260925", "smlt_options_20260925.csv"), ("Covariance_OOS", "covariance_oos.csv"),
              ("Smart_comparison", "smart_portfolio_comparison.csv"), ("Robust_solutions", "robust_optimizer_solutions.csv"),
              ("Collateral_conclusions", "collateral_conclusions.csv"), ("Factor_decomposition", "factor_decomposition.csv"),
              ("Futures_spec_audit", "futures_contract_spec_audit.csv"), ("Complexity_multidim", "complexity_multidim.csv"),
              ("Liquidity_waterfall", "liquidity_waterfall_results.csv"), ("Risk_contributions", "risk_contributions.csv")]
    for title, fn in sheets:
        p = ctx.out / fn
        if not p.exists():
            continue
        df = pd.read_csv(p)
        w = wb.create_sheet(title)
        for r in dataframe_to_rows(df, index=False, header=True):
            w.append([None if (isinstance(v, float) and np.isnan(v)) else v for v in r])
        for c in w[1]:
            c.font = Font(bold=True)
        w.freeze_panes = "A2"
    # live calculator
    mi = pd.read_csv(ctx.out / "market_inputs_at_cutoff.csv", index_col=0)["value"]
    c = wb.create_sheet("Calculator", 1)
    fill = PatternFill("solid", fgColor="DDEBF7")
    rows = [("E: SMLT exposure, RUB", ctx.E), ("C: hedge capital, RUB", ctx.C), ("S0: SMLT close 2026-09-25", float(mi["S0"])),
            ("h: futures hedge ratio (share units)", 0.75), ("x: SMLT terminal return", -0.25),
            ("tau: horizon, years", 0.5), ("b: futures basis, annualised", float(mi["basis_ann"])),
            ("y: money-market yield, annual", float(mi["mmf_yield_LQDT_60d"])), ("mr: NCC IM rate", float(mi["mr1"])),
            ("cost: one-way futures cost, fraction of notional", float(mi["fut_cost_rate_one_way"])),
            ("d: share of SMLT sold (DE_RISKING)", 0.0)]
    c["A1"], c["B1"] = "Input", "Value"
    for i, (k, v) in enumerate(rows, 2):
        c.cell(row=i, column=1, value=k)
        c.cell(row=i, column=2, value=v).fill = fill
    f = [("Shares held", "=B2*(1-B12)/B4"), ("Futures contracts to sell (lot = 1 share)", "=ROUND(B5*B15,0)"),
         ("Initial margin, RUB", "=B10*B5*B2*(1-B12)"), ("IM as % of hedge capital", "=B17/B3"),
         ("Spot P&L", "=B2*(1-B12)*B6"), ("Futures P&L incl. basis", "=-B5*B2*(1-B12)*B6+B5*B2*(1-B12)*B8*B7"),
         ("Futures costs (entry + exit + 2 legs per roll)", "=B5*B2*(1-B12)*B11*(2+2*ROUNDUP(B7/0.25,0))"),
         ("Sleeve interest (IM earns 0)", "=MAX(B3+B2*B12-B17,0)*B9*B7"),
         ("Total P&L, RUB", "=B19+B20-B21+B22-B2*B12*0.001"), ("Total P&L, % of RUB 10m SMLT", "=B23/B2"),
         ("Total P&L, % of total capital", "=B23/(B2+B3)"), ("Day-1 VM if SMLT jumps +20%, RUB", "=-B5*B2*(1-B12)*0.2")]
    for i, (k, formula) in enumerate(f, 15):
        c.cell(row=i, column=1, value=k)
        c.cell(row=i, column=2, value=formula)
    c.column_dimensions["A"].width = 52
    c.column_dimensions["B"].width = 18
    import datetime as _dt
    fixed = _dt.datetime(2026, 9, 25, 0, 0, 0)            # deterministic file: no wall-clock timestamps
    wb.properties.created = fixed
    wb.properties.modified = fixed
    wb.properties.lastModifiedBy = "smlt_hedge pipeline"
    wb.save(ctx.out / "SMLT_Hedging_Model.xlsx")
    _normalise_xlsx(ctx.out / "SMLT_Hedging_Model.xlsx")


def _normalise_xlsx(path) -> None:
    """openpyxl stamps save-time into docProps/core.xml and zip entry dates -> rewrite both deterministically."""
    import re
    import zipfile
    items = []
    with zipfile.ZipFile(path) as z:
        for info in z.infolist():
            b = z.read(info.filename)
            if info.filename == "docProps/core.xml":
                b = re.sub(rb"<dcterms:modified([^>]*)>[^<]*</dcterms:modified>",
                           rb"<dcterms:modified\1>2026-09-25T00:00:00Z</dcterms:modified>", b)
                b = re.sub(rb"<dcterms:created([^>]*)>[^<]*</dcterms:created>",
                           rb"<dcterms:created\1>2026-09-25T00:00:00Z</dcterms:created>", b)
            items.append((info.filename, b))
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as z:
        for name, b in items:
            zi = zipfile.ZipInfo(name, date_time=(2026, 9, 25, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, b)


def s10_report(ctx) -> None:
    from .tables import build_tables
    build_tables(ctx)
    for f in (fig_smlt, fig_credit, fig_rolling, fig_im, fig_capital, fig_complexity, fig_tradeoff, fig_lar,
              fig_surface, fig_walkforward):
        try:
            f(ctx)
        except Exception as e:  # noqa: BLE001  one failing figure must not block the others; logged
            ctx.log.setdefault("figure_errors", {})[f.__name__] = repr(e)
    from . import figures2
    if (ctx.out / "smart_portfolio_comparison.csv").exists():
        figures2.all_figures(ctx)
    excel(ctx)
    from . import docs, dossier, findings, repair_report
    docs.write_docs(ctx)
    dossier.write(ctx)
    findings.render(ctx)
    if (ctx.out / "baseline_v1" / "key_findings.json").exists() and (ctx.out / "collateral_conclusions.csv").exists():
        repair_report.write(ctx)
